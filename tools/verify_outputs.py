"""Independent, optional QA: pandas source recomputation plus XLSX inspection."""
import hashlib,json,sqlite3,zipfile
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
s=json.loads((ROOT/'results/analysis.json').read_text(encoding='utf-8'))
d=pd.read_csv(ROOT/'data/raw/data.csv',encoding='latin1',dtype=str,keep_default_na=False)
q=pd.to_numeric(d.Quantity,errors='raise'); price=pd.to_numeric(d.UnitPrice,errors='raise')
pd.to_datetime(d.InvoiceDate,format='%m/%d/%Y %H:%M',errors='raise')
amount=pd.Series([int((Decimal(p)*Decimal(n)*100).quantize(Decimal(1),rounding=ROUND_HALF_UP)) for p,n in zip(d.UnitPrice,d.Quantity)])
sale=(price>0)&(q>0)&~d.InvoiceNo.str.upper().str.startswith('C')
cancel=(price>0)&(q<0)&d.InvoiceNo.str.upper().str.startswith('C')
eligible=sale|cancel; repeated=d.duplicated(keep='first')
checks={
 'source_hash':hashlib.sha256((ROOT/'data/raw/data.csv').read_bytes()).hexdigest()==s['input_sha256'],
 'rows':len(d)==s['metrics']['input_rows'],
 'repeats':int(repeated.sum())==s['metrics']['exact_repeat_rows'],
 'missing_customer':int((d.CustomerID=='').sum())==s['metrics']['missing_customer_rows'],
 'sale_rows':int(sale.sum())==s['metrics']['sale'],
 'cancellation_rows':int(cancel.sum())==s['metrics']['cancellation'],
 'excluded_rows':int((~eligible).sum())==s['metrics']['excluded'],
 'net_pence':int(amount[eligible].sum())==s['metrics']['net_pence'],
 'deduplicated_net_pence':int(amount[eligible&~repeated].sum())==s['metrics']['deduplicated_net_pence'],
 'gross_sales_pence':int(amount[sale].sum())==s['metrics']['gross_sales_pence'],
 'unidentified_gross':int(amount[sale&(d.CustomerID=='')].sum())==s['metrics']['unidentified_gross_pence'],
}
with sqlite3.connect(ROOT/'results/retail.sqlite') as db:
    checks['sqlite_integrity']=db.execute('PRAGMA integrity_check').fetchone()[0]=='ok'
    checks['no_orphan_lines']=db.execute('SELECT COUNT(*) FROM lines l LEFT JOIN raw_rows r USING(row_id) WHERE r.row_id IS NULL').fetchone()[0]==0
book=ROOT/'deliverables/retail-sales-review.xlsx'
wb=load_workbook(book,data_only=True); wf=load_workbook(book,data_only=False)
checks['xlsx_cached_net']=abs(wb['Sales review']['B8'].value-s['metrics']['net_pence']/100)<.005
checks['xlsx_net_formula']=wf['Sales review']['B8'].data_type=='f'
checks['xlsx_chart']=len(wf['Sales review']._charts)==1
checks['xlsx_country_shares']=abs(sum(wb['Country detail'].cell(r,6).value for r in range(6,44))-1)<1e-10
with zipfile.ZipFile(book) as archive:
    chart=archive.read(next(n for n in archive.namelist() if n.endswith('/chart1.xml'))).decode()
    checks['xlsx_chart_line_color']='3067B0' in chart.upper()
    checks['xlsx_chart_formula_binding']='$D$22:$D$34' in chart
result={'method':'Independent pandas source recomputation; read-only openpyxl/XML checks; SQLite integrity check',
        'checks':checks,'all_passed':all(checks.values()),'nativeExcelRecalculationTested':False}
(ROOT/'docs/independent-verification.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
print(json.dumps(result,indent=2))
if not all(checks.values()): raise SystemExit(1)
